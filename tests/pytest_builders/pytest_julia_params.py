"""
Julia builder (ODE backend): changing a model's inputs between runs.

A generated model's ``params(; name = value, ..., sheets, files)`` gives its
constants, lookups and data other values, and what is computed from them
(other constants, the initial values of the stocks) follows.  Each scenario
below is run in Julia and compared, at every saved time, with Python PySD
run with the same change (``model.run(params=...)``, or the model with the
other Excel sheet or file written in).

All scenarios run in one Julia process.  Environment variables as in
``pytest_julia_reference.py`` (``PYSD_JULIA_REQUIRED``).
"""
import os
import subprocess
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import xarray as xr

from tests.pytest_builders.pytest_julia_reference import _julia_unavailable

pytestmark = pytest.mark.julia

MODEL = r"""{UTF-8}
region:
	north, south, east
	~
	~		|

growth rate=
	0.05
	~	1/Month
	~		|

initial population=
	100
	~
	~		|

doubled=
	initial population * 2
	~
	~		|

population= INTEG (
	births,
		doubled)
	~
	~		|

births=
	population * growth rate * effect(Time) * scale
	~
	~		|

effect(
	[(0,0)-(10,2)],(0,1),(10,2))
	~
	~		|

weights[region]=
	1, 2, 3
	~
	~		|

share[region]=
	weights[region] / SUM(weights[region!])
	~
	~		|

by region[region]=
	population * share[region]
	~
	~		|

scale=
	GET DIRECT CONSTANTS('inputs.xlsx', 'base', 'B2')
	~
	~		|

policy=
	GET DIRECT LOOKUPS('inputs.xlsx', 'base', '4', 'B5')
	~
	~		|

policy now=
	policy(Time)
	~
	~		|

regional[region]=
	GET DIRECT DATA('inputs.xlsx', 'base', '4', 'B7')
	~
	~		|

regional total=
	SUM(regional[region!])
	~
	~		|

initial births=
	INITIAL(births)
	~
	~		|

initial growth=
	INITIAL(growth rate * 10)
	~
	~		|

delay time=
	2
	~	Month
	~		|

delayed births=
	DELAY FIXED(births, delay time, initial growth)
	~
	~		|

safe log=
	IF THEN ELSE(growth rate < 1, 0, LN(growth rate - 1))
	~
	~	LN of a negative number in the branch that is not taken.	|

safe root=
	IF THEN ELSE(growth rate < 1, 1, SQRT(growth rate - 1))
	~
	~		|

FINAL TIME  = 10
	~	Month
	~		|

INITIAL TIME  = 0
	~	Month
	~		|

SAVEPER  = 1
	~	Month
	~		|

TIME STEP  = 0.25
	~	Month
	~		|
"""

#: Python column -> Julia name and index in observe's output
COLUMNS = {
    "population": ("population", None),
    "policy now": ("policy_now", None),
    "regional total": ("regional_total", None),
    "by region[north]": ("by_region", 1),
    "by region[east]": ("by_region", 3),
    "share[south]": ("share", 2),
    "regional[south]": ("regional", 2),
    "initial births": ("initial_births", None),
    "initial growth": ("initial_growth", None),
    "delayed births": ("delayed_births", None),
    # LN and SQRT out of their domain give NaN, as numpy (not an error)
    "safe log": ("safe_log", None),
    "safe root": ("safe_root", None),
}

REGIONS = {"region": ["north", "south", "east"]}

#: scenario -> (Julia params(...) arguments, Python run params, model change)
SCENARIOS = {
    "base": ("", {}, None),
    "constant": ("growth_rate = 0.1", {"growth rate": 0.1}, None),
    # the initial value of the stock follows the constant
    "initial value": ("initial_population = 50", {"initial population": 50}, None),
    # a constant computed from others can be given too
    "computed constant": ("doubled = 10", {"doubled": 10}, None),
    "array": ("weights = [3, 2, 1]",
              {"weights": xr.DataArray([3, 2, 1], REGIONS, ["region"])}, None),
    "array from a number": ("weights = 2",
                            {"weights": xr.DataArray([2, 2, 2], REGIONS, ["region"])}, None),
    "external constant": ("scale = 2", {"scale": 2}, None),
    "lookup as series": ("effect = ([0, 10], [1, 1])",
                         {"effect": pd.Series([1.0, 1.0], index=[0.0, 10.0])}, None),
    "lookup as function": ("effect = x -> 1.0",
                           {"effect": pd.Series([1.0, 1.0], index=[0.0, 10.0])}, None),
    "subscripted data": (
        "regional = ([0, 10], [1 2 3; 7 8 9])",
        {"regional": pd.Series(
            [xr.DataArray([1.0, 2, 3], REGIONS, ["region"]),
             xr.DataArray([7.0, 8, 9], REGIONS, ["region"])], index=[0.0, 10.0])},
        None),
    "several at once": ("growth_rate = 0.1, initial_population = 50, scale = 2",
                        {"growth rate": 0.1, "initial population": 50, "scale": 2}, None),
    "sheet": ('sheets = Dict("base" => "high")', {}, ("'base'", "'high'")),
    "file": ('files = Dict("inputs.xlsx" => "other.xlsx")', {},
             ("'inputs.xlsx'", "'other.xlsx'")),
    # the defaults are untouched by the runs above
    "base again": ("", {}, None),
}

#: wrong input -> (Julia params(...) arguments, exception, text in the message)
ERRORS = {
    "unknown name": ("no_such_input = 1", "ArgumentError", "no_such_input"),
    "variable of the run": ("births = 1", "ArgumentError", "births"),
    "array of another size": ("weights = [1, 2]", "DimensionMismatch", "size (3,)"),
    "array for a number": ("growth_rate = [1, 2]", "ArgumentError", "growth_rate"),
    "not a lookup": ('effect = "x"', "ArgumentError", "lookup"),
    # the number of stages of a DELAY FIXED is set when the model is translated
    "fixed by the translation": ("delay_time = 3", "ArgumentError",
                                 "delay_time sets the number of stages of DELAY FIXED"),
}


def _workbook(path, sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (scale, policy, regional) in sheets.items():
        ws = wb.create_sheet(name)
        ws["B2"] = scale
        for k, x in enumerate([0, 5, 10]):
            ws.cell(row=4, column=2 + k, value=x)
            ws.cell(row=5, column=2 + k, value=policy[k])
            for r in range(3):
                ws.cell(row=7 + r, column=2 + k, value=regional * (r + 1) + k)
    wb.save(path)


def _julia_script(model_jl):
    runs = "\n".join(
        f'push!(runs, "{name}" => () -> params({args}))'
        for name, (args, _, _) in SCENARIOS.items())
    errors = "\n".join(
        f'push!(errors, "{name}" => () -> params({args}))'
        for name, (args, _, _) in ERRORS.items())
    columns = ", ".join(f'("{jl}", {idx if idx else 0})' for jl, idx in COLUMNS.values())
    # one line per result: kind|scenario|comma-separated numbers or a message
    return f'''
include("{model_jl.as_posix()}")
columns = [{columns}]
row(o) = [k == 0 ? Float64(o[n]) : Float64(o[n][k]) for (n, k) in columns]
runs, errors = Pair{{String,Function}}[], Pair{{String,Function}}[]
{runs}
{errors}
for (name, make) in runs
    p = make()
    sol = run_model(params = p, saveat = initial_time:saveper:final_time)
    println("t|", name, "|", join(repr.(sol.t), ","))
    for (t, u) in zip(sol.t, sol.u)
        println("row|", name, "|", join(repr.(row(observe(u, t; p = p))), ","))
    end
    println("initial|", name, "|", join(repr.(initial_state(p)), ","))
end
for (name, make) in errors
    msg = try
        make(); "no error"
    catch e
        string(nameof(typeof(e))) * ": " * sprint(showerror, e)
    end
    println("error|", name, "|", replace(msg, "\\n" => " "))
end
'''


def _parse(stdout):
    out = {}
    for line in stdout.splitlines():
        kind, sep, rest = line.partition("|")
        if not sep or kind not in ("t", "row", "initial", "error"):
            continue
        name, _, data = rest.partition("|")
        entry = out.setdefault(name, {"values": []})
        if kind == "error":
            entry["error"] = data
        elif kind == "row":
            entry["values"].append([float(x) for x in data.split(",")])
        else:
            entry[kind] = [float(x) for x in data.split(",") if x]
    return out


@pytest.fixture(scope="module")
def results(tmp_path_factory):
    reason = _julia_unavailable("ode")
    if reason:
        if os.environ.get("PYSD_JULIA_REQUIRED"):
            pytest.fail(reason)
        pytest.skip(reason)
    import pysd
    from pysd import translate_to_julia

    folder = tmp_path_factory.mktemp("julia_params")
    mdl = folder / "scenarios.mdl"
    mdl.write_text(MODEL)
    _workbook(folder / "inputs.xlsx", {"base": (1.0, [0, 1, 2], 1),
                                       "high": (1.5, [1, 3, 5], 10)})
    _workbook(folder / "other.xlsx", {"base": (0.5, [2, 2, 2], 4)})
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        model_jl = translate_to_julia(mdl)

    python = {}
    for name, (_, params, change) in SCENARIOS.items():
        path = mdl
        if change:
            path = folder / f"scenarios_{name}.mdl"
            path.write_text(MODEL.replace(*change))
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            out = pysd.read_vensim(path).run(params=params,
                                             return_columns=list(COLUMNS))
        python[name] = out

    script = folder / "run_scenarios.jl"
    script.write_text(_julia_script(model_jl))
    run = subprocess.run(["julia", "--startup-file=no", str(script)],
                         capture_output=True, text=True, timeout=3600)
    assert run.returncode == 0, run.stdout[-3000:] + run.stderr[-3000:]
    return python, _parse(run.stdout)


@pytest.mark.parametrize("scenario", [n for n in SCENARIOS])
def test_scenario_matches_python(results, scenario):
    python, julia = results
    jl = julia[scenario]
    py = python[scenario]
    assert np.allclose(jl["t"], py.index.values)
    np.testing.assert_allclose(np.array(jl["values"]), py[list(COLUMNS)].values,
                               rtol=1e-10, atol=1e-12)


def test_scenarios_change_the_results(results):
    """Each scenario changes something (or the comparison above proves little)."""
    _, julia = results
    base = np.array(julia["base"]["values"])
    for name in SCENARIOS:
        if name in ("base", "base again"):
            continue
        assert not np.allclose(np.array(julia[name]["values"]), base), name


def test_initial_state_follows_the_constants(results):
    _, julia = results
    # population starts at doubled = 2 * initial population
    assert julia["base"]["initial"][0] == 200.0
    assert julia["initial value"]["initial"][0] == 100.0
    assert julia["computed constant"]["initial"][0] == 10.0


@pytest.mark.parametrize("case", [n for n in ERRORS])
def test_wrong_inputs_are_errors(results, case):
    _, julia = results
    _, exc, text = ERRORS[case]
    message = julia[case]["error"]
    assert message.startswith(exc + ": "), message
    assert text in message, message
